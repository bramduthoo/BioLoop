# -*- coding: utf-8 -*-
"""Apply the 2026-09-08 selection-control review round to the `Streams` sheet.

A one-off. Kept in `migrations/` for provenance - do NOT re-run it as part of the pipeline
(`migrations/README.md`). It is idempotent, so a second run reports "already applied" rather
than doubling anything, but the finished state is what the workbook carries.

What it does, and why - the reviewer's decisions, each traced to a source page:

A1  EIEREN TO L4.  `commodity_hierarchy.md` line 25: "a single crop, species or product is never
    an L3". Every egg row in the corpus sat at L3 with L4 blank, so eggs were invisible to a
    selection, which only reaches L4/L5. `Dierlijk - vee / Eieren` was the only bucket-D group
    naming a product.

A2  SPLIT THE DAIRY NODE.  OVAM Tabel 30 (S002 p.54) is titled "Bestemmingen van
    voedselreststromen in de voedingsindustrie" and its rows are the eight food-industry
    sub-sectors summing to 1.999.383 t; `Zuivel` (123.219) is one of them, beside Bakkerij and
    Dranken. MONBIO S091 p.94 says the same in words: "De grootste productie in deze sector is
    natuurlijk de melk zelf. Typische nevenstromen en productieresiduen in deze sector zijn
    melkwei en behandeld zuiveringsslib." C-100/C-200 were nevertheless placed on L4 = Melk, so
    the node added 12.378 t of farm milk to 123.219 t of factory residue. They move to their own
    L4; the remaining `Melk` rows now mean only milk.

A3  CAPTURE MELKWEI (C-802, C-803).  S091 p.168 Tabel 65 and S007 p.160 Tabel 47 both print
    "NACE 10.5 zuivelfabrieken -> 70.167 totaal, waarvan melkwei 49.722 en zuiveringsslib
    20.445". The split was skipped under the cross-source restatement rule because S005 owns the
    2018 estimate - verified correct, S005 p.131 carries the identical row - but S005 is
    _RETIRED, so the figure landed nowhere while state.md recorded whey as absent from all 801
    claims. Reviewer decision: capture it as a narrow, documented exception.
    THE EXCEPTION IS THE ZUIVEL CELL ONLY, NOT TABEL 65. The rest of that table is genuinely
    S005's; re-capturing it would duplicate ~1,8 Mt.

A4  NORMALISE THREE QUANTITY-TYPE TOTALS.  C-154/C-157/C-158 each equal exactly
    voedselverlies + nevenstroom, which is the definition of agri-food waste, not evidence of
    aggregation. They carried the prefix only because OVAM 2020 wrote "(totaal)"; the identical
    2023 rows (C-041/C-046/C-047) did not and are live claims. `promote_totals.py` now knows the
    difference, so this does not come back.

A5  DECLARE THE EXCLUDED COMPONENT.  Protocol v2.3: an aggregate that includes a deliberately
    excluded component must say so on the row. 20.445 of C-282/C-467's 70.167 t is
    zuiveringsslib, which `quantity_type.csv` puts out of scope.

A6  SURFACE THE HIDES CAVEAT.  Protocol v2.3: a caveat that changes what a figure measures
    belongs in `stream_name_NL`. C-277/C-462 count 100.000 'stuks' huiden as if they were
    tonnes; that sat only in `source_type_label`, so it never reached the deliverable.

    database/.venv/Scripts/python database/register/migrations/apply_review_2026-09-08.py --dry-run
"""
import pathlib, sys
import openpyxl

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent                      # register/ - HERE is register/migrations/
WB = ROOT / "BIOLOOP_streams_and_sources.xlsx"
SHEET = "Streams"
DRY = "--dry-run" in sys.argv

# ---------------------------------------------------------------------------------------------
# A1 - eggs to L4.  The six residual rows are what makes eggs selectable; the two plain
# production rows follow so `level_1to5` stays the deepest filled column on every egg row.
# C-375/C-568 (Vogeleieren uit de schaal) are deliberately NOT touched - a processed product,
# open item O3.
EGGS = ["C-047", "C-060", "C-071", "C-158", "C-170", "C-182", "C-262", "C-443"]

# A2 - the dairy sector rows, with the caveat moved into the name
ZUIVEL = {
    "C-100": "Voedselreststromen Zuivel (NACE-sectorrij voedingsindustrie - fabrieksresidu, "
             "niet melk als product)",
    "C-200": "Voedselreststromen voedingsindustrie - Zuivel (NACE-sectorrij - fabrieksresidu, "
             "niet melk als product)",
}
ZUIVEL_ALT = ("zusterrijen van dezelfde sectortabel: Bakkerij, Aardappelen/groenten/fruit, "
              "Dranken, Olien-vetten, Suiker-chocolade, Deegwaren-zetmeel, Vlees-vis-gevogelte "
              "(alle boven L4 gelaten)")

# A4 - strip the prefix from the quantity-type totals
DEPROMOTE = ["C-154", "C-157", "C-158"]

# A5 / A6 - caveats that belong in the name
NAME_SUFFIX = {
    "C-282": " (incl. 20.445 t zuiveringsslib, buiten scope - zie C-802 voor de wei)",
    "C-467": " (incl. 20.445 t zuiveringsslib, buiten scope - zie C-803 voor de wei)",
    "C-277": " (incl. 100.000 'stuks' huiden die de bron als ton meetelt: 531.159 + 100.000)",
    "C-462": " (incl. 100.000 'stuks' huiden die de bron als ton meetelt: 494.819 + 100.000)",
}

# A3 - the two new whey claims
NEW = [
    {"claim_id": "C-802", "source_id": "S091", "source_short": "MONBIO 4.0",
     "source_page": "168", "source_table_figure": "Tabel 65",
     "also_stated_in": "sectortotaal 70.167 t in dezelfde tabelrij - zie C-282; tekst p.94; "
                       "Prodcom 105155 'Wei' staat in Tabel 35 als vertrouwelijk (C)"},
    {"claim_id": "C-803", "source_id": "S007", "source_short": "MONBIO 3.0",
     "source_page": "160", "source_table_figure": "Tabel 47",
     "also_stated_in": "sectortotaal 70.167 t in dezelfde tabelrij - zie C-467; tekst p.118; "
                       "Prodcom 105155 'Wei' staat vertrouwelijk (C)"},
]
NEW_SHARED = {
    "stream_name_NL": "Melkwei (nevenstroom zuivelsector; OVAM/IMJV-schatting 2018, door de bron "
                      "overgenomen)",
    "L1_role": "Reststroom", "L2_commodity_group": "Dierlijk - vee",
    "L3_commodity_subgroup": "Melk", "L4_ingredient": "Zuivelnevenstroom",
    "L5_fraction_as_named": None, "level_1to5": 4,
    "chain_stage": "Vervaardiging van zuivelproducten", "chain_L2": "Voedingsindustrie",
    "reference_year": "2018", "volume_t_per_yr": 49722, "value_as_reported": 49722,
    "unit_as_reported": "ton", "conversion_factor_to_t_per_yr": 1,
    "quantity_type": "agri-food waste",
    "source_type_label": "nevenstromen en productieresiduen zonder afvalstatuut, NACE 10.5 "
                         "zuivelfabrieken en kaasmakerijen; de sectorrij telt 70.167 t = melkwei "
                         "49.722 + zuiveringsslib 20.445, waarvan het slib buiten scope valt "
                         "(quantity_type.csv). MONBIO's residuvocabulaire is economisch, dus "
                         "agri-food waste met type_assumed",
    "geography": "Vlaanderen", "provenance": "read in PDF", "DECISION_expert": None,
}


def main():
    if not WB.exists():
        sys.exit(f"missing {WB.name}")
    wb = openpyxl.load_workbook(WB)
    ws = wb[SHEET]
    col = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    need = set(NEW_SHARED) | {"claim_id", "type_assumed", "source_page", "source_table_figure",
                              "also_stated_in", "source_id", "source_short"}
    missing = [c for c in need if c not in col]
    if missing:
        sys.exit(f"columns not found in {SHEET}: {missing}")

    at = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=col["claim_id"]).value
        if v:
            at[str(v).strip()] = r

    def get(cid, name):
        return ws.cell(row=at[cid], column=col[name]).value

    def put(cid, name, value):
        ws.cell(row=at[cid], column=col[name]).value = value

    done = []

    # -- A1 ------------------------------------------------------------------------------
    for cid in EGGS:
        if cid not in at:
            sys.exit(f"A1: {cid} not found")
        if (get(cid, "L4_ingredient") or "") != "Eieren":
            put(cid, "L4_ingredient", "Eieren")
            put(cid, "level_1to5", 4)
            done.append(f"A1  {cid}  L4 = Eieren, level 4")

    # -- A2 ------------------------------------------------------------------------------
    for cid, newname in ZUIVEL.items():
        if (get(cid, "L4_ingredient") or "") != "Zuivelnevenstroom":
            put(cid, "L4_ingredient", "Zuivelnevenstroom")
            put(cid, "level_1to5", 4)
            put(cid, "stream_name_NL", newname)
            if not (get(cid, "also_stated_in") or "").strip():
                put(cid, "also_stated_in", ZUIVEL_ALT)
            done.append(f"A2  {cid}  L4 = Zuivelnevenstroom, level 4, naam + zusterrijen")

    # -- A4 ------------------------------------------------------------------------------
    for cid in DEPROMOTE:
        name = str(get(cid, "stream_name_NL") or "")
        if name.startswith("AGGREGAAT - "):
            put(cid, "stream_name_NL", name[len("AGGREGAAT - "):])
            done.append(f"A4  {cid}  prefix verwijderd -> {name[len('AGGREGAAT - '):][:52]}")

    # -- A5 / A6 -------------------------------------------------------------------------
    for cid, suffix in NAME_SUFFIX.items():
        name = str(get(cid, "stream_name_NL") or "")
        if suffix.strip() not in name:
            put(cid, "stream_name_NL", name + suffix)
            done.append(f"A5/A6  {cid}  caveat in de naam gezet")
    # C-277 carried a stray English note in `also_stated_in`, a column for restatement locations
    stray = get("C-277", "also_stated_in")
    if stray and "as if they were tonnes" in str(stray):
        put("C-277", "also_stated_in", None)
        done.append("A6  C-277  losse notitie uit also_stated_in gehaald (staat nu in de naam)")

    # -- A3 ------------------------------------------------------------------------------
    proto = at.get("C-282")
    for spec in NEW:
        cid = spec["claim_id"]
        if cid in at:
            continue
        r = ws.max_row + 1
        for k, v in list(NEW_SHARED.items()) + list(spec.items()):
            ws.cell(row=r, column=col[k]).value = v
        # copy the boolean's stored form and the factor's number format from an existing row,
        # so Excel renders them exactly like their neighbours (v2.3: the factor must stay legible)
        ws.cell(row=r, column=col["type_assumed"]).value = \
            ws.cell(row=proto, column=col["type_assumed"]).value
        ws.cell(row=r, column=col["conversion_factor_to_t_per_yr"]).number_format = \
            ws.cell(row=proto, column=col["conversion_factor_to_t_per_yr"]).number_format
        at[cid] = r
        done.append(f"A3  {cid}  nieuwe claim, 49.722 t melkwei uit {spec['source_short']} "
                    f"{spec['source_table_figure']} (p.{spec['source_page']})")

    if not done:
        print("already applied - nothing to change")
        return
    print(f"{len(done)} wijziging(en):\n")
    for d in done:
        print("  " + d)
    if DRY:
        print("\n--dry-run: workbook not written")
        return
    try:
        wb.save(WB)
    except PermissionError:
        sys.exit(f"\nCannot write {WB.name} - it is open in Excel. Close it and run this again.")
    print(f"\nwrote {WB.name} - now re-run the pipeline from prep_data.py")


if __name__ == "__main__":
    main()
