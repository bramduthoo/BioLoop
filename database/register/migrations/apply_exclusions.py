"""Write the reviewer's exclusions into `DECISION_expert` in the register workbook.

During the varia_reclass review the reviewer marked 44 claims with a remark saying the claim
itself should leave the register ("excluded from original excel file, so we can ignore these").
That decision belongs in the workbook's own `DECISION_expert` column, not in a crosswalk: the
crosswalk decides *placement*, the workbook decides *whether the claim counts at all*.

Reads the crosswalk, finds every row whose remark says "exclude"/"ignore", and writes
`DECISION_expert` for those claims. Never overwrites a value already there - an existing
expert decision always wins, and is reported instead.

Reads and writes the `Streams` sheet BY COLUMN HEADER. Idempotent. Writes the workbook last,
and says so plainly if Excel is holding the file open.

    database/.venv/Scripts/python.exe database/register/apply_exclusions.py --dry-run
    database/.venv/Scripts/python.exe database/register/apply_exclusions.py
"""
import csv, io, pathlib, re, sys
import openpyxl

HERE = pathlib.Path(__file__).resolve().parent
WB = HERE / "BIOLOOP_streams_and_sources.xlsx"
CROSSWALK = HERE / "crosswalks" / "varia_reclass.csv"
SHEET = "Streams"
DRY = "--dry-run" in sys.argv
STAMP = "no - reviewer excluded this claim (varia_reclass review, 2026-08-31)"

# a remark counts as an exclusion instruction only if it says so; "maybe an aggregate" does not.
# Matches on "clude" so the reviewer's "xcluded" typo is caught alongside "excluded".
EXCL_RE = re.compile(r"(clude|ignore)", re.I)
KEEP_RE = re.compile(r"aggregat|subset|L3|L4", re.I)


def remark_of(row):
    """The reviewer's free-text remark: the unnamed trailing column, or `remark`/`opmerking`."""
    for k in ("", "remark", "opmerking", "note_reviewer"):
        if k in row and (row.get(k) or "").strip():
            return row[k].strip()
    return ""


def main():
    if not CROSSWALK.exists():
        sys.exit(f"missing {CROSSWALK.name}")
    if not WB.exists():
        sys.exit(f"missing {WB.name}")

    rows = list(csv.DictReader(io.open(CROSSWALK, encoding="utf-8-sig"), delimiter=";"))
    want = {}
    for r in rows:
        rem = remark_of(r)
        decision = (r.get("DECISION") or "").strip().lower()
        if decision in ("no", "skip", "exclude") and EXCL_RE.search(rem) and not KEEP_RE.search(rem):
            want[r["claim_id"]] = rem
    print(f"{len(want)} claim(s) flagged for exclusion by a reviewer remark")

    wb = openpyxl.load_workbook(WB)
    ws = wb[SHEET]
    col = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    for need in ("claim_id", "DECISION_expert", "stream_name_NL"):
        if need not in col:
            sys.exit(f"column '{need}' not found in {SHEET}")

    wrote, already, conflict, missing = 0, 0, [], set(want)
    for row in range(2, ws.max_row + 1):
        cid = ws.cell(row=row, column=col["claim_id"]).value
        cid = str(cid).strip() if cid else ""
        if cid not in want:
            continue
        missing.discard(cid)
        cell = ws.cell(row=row, column=col["DECISION_expert"])
        cur = "" if cell.value is None else str(cell.value).strip()
        if not cur:
            cell.value = STAMP
            wrote += 1
        elif cur.lower().split()[0].strip(":-,.") in ("no", "nee", "exclude"):
            already += 1
        else:
            conflict.append((cid, cur))

    print(f"  {wrote} newly marked · {already} already excluded")
    if conflict:
        print(f"  ! {len(conflict)} already carry a different expert decision - left untouched:")
        for cid, cur in conflict:
            print(f"      {cid}: {cur[:70]}")
    if missing:
        print(f"  ! not found in the sheet: {', '.join(sorted(missing))}")

    if DRY:
        print("--dry-run: workbook not written")
        return
    if not wrote:
        print("nothing to write")
        return
    try:
        wb.save(WB)
    except PermissionError:
        sys.exit(f"\nCannot write {WB.name} - it is open in Excel. Close it and run this again.")
    print(f"wrote {WB.name}")


if __name__ == "__main__":
    main()
